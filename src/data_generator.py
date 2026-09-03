"""
TalentMatch ML - Data Generation Module
Generates a controlled, reproducible prototype candidate corpus across PDF, DOCX, and TXT formats,
and structured Job Descriptions across Machine Learning, Full Stack, and DevOps roles.
"""

import os
import json
import docx

def generate_sample_resumes(resumes_dir: str):
    """Creates synthetic candidate resumes across varying seniorities and tech stacks."""
    os.makedirs(resumes_dir, exist_ok=True)
    
    # 1. Lead ML Engineer (.docx)
    doc1 = docx.Document()
    doc1.add_heading("Alex Rivera — Senior Machine Learning & MLOps Engineer", 0)
    doc1.add_paragraph("Email: alex.rivera.ml@example.com | Phone: (555) 234-5678 | GitHub: github.com/arivera-ml")
    doc1.add_heading("Professional Summary", level=1)
    doc1.add_paragraph("Senior Machine Learning Engineer with 6.5 years of experience building, training, and deploying large-scale deep learning models, LLM pipelines, and automated MLOps infrastructure in cloud environments.")
    doc1.add_heading("Core Technical Competencies", level=1)
    doc1.add_paragraph("Languages & Frameworks: Python, PyTorch, TensorFlow, Scikit-learn, XGBoost, LightGBM, HuggingFace, Transformers, LangChain, RAG.")
    doc1.add_paragraph("Data Engineering & DBs: PostgreSQL, Redis, Apache Spark, SQL, Data Pipelines, ETL.")
    doc1.add_paragraph("Cloud & MLOps: AWS, Docker, Kubernetes, CI/CD, MLflow, Model Deployment, Feature Engineering, Linux.")
    doc1.add_heading("Experience", level=1)
    doc1.add_paragraph("Lead ML Engineer at NeuralScale Inc (2021 - Present): Architected end-to-end NLP retrieval-augmented generation pipelines using HuggingFace and PyTorch. Deployed models to Kubernetes clusters on AWS serving 10M requests daily.")
    doc1.add_paragraph("Machine Learning Engineer at DataVibe (2018 - 2021): Built time series forecasting and gradient boosting models using LightGBM and Scikit-learn.")
    doc1.add_heading("Education", level=1)
    doc1.add_paragraph("Master's Degree in Computer Science, Stanford University (2018)")
    doc1.save(os.path.join(resumes_dir, "candidate_01_lead_ml_engineer.docx"))

    # 2. Mid-Level Data Scientist (.txt)
    txt2 = """Sarah Chen — Data Scientist & Quantitative Analyst
Email: schen.ds@gmail.com | Phone: +1-555-876-5432 | LinkedIn: linkedin.com/in/sarahchen-ds

Professional Summary:
Data Scientist with 4.0 years of experience specializing in predictive modeling, statistical analysis, feature engineering, and time series forecasting.

Technical Skills:
- Programming: Python, SQL, R
- Machine Learning: Scikit-learn, XGBoost, LightGBM, Random Forest, Time Series Forecasting, A/B Testing
- Data Analysis & Tools: Pandas, NumPy, Matplotlib, Seaborn, Tableau, Jupyter
- Databases & Cloud: PostgreSQL, BigQuery, AWS, Git

Experience:
Data Scientist at RetailMetrics (2020 - Present):
- Developed customer churn and weekly demand forecasting models using XGBoost and Pandas.
- Designed automated A/B testing frameworks improving marketing conversion by 14%.

Education:
Master's Degree in Applied Statistics, University of Michigan (2020)
"""
    with open(os.path.join(resumes_dir, "candidate_02_data_scientist.txt"), "w", encoding="utf-8") as f:
        f.write(txt2)

    # 3. Junior ML Developer (.pdf)
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.pdfgen import canvas
        pdf_path = os.path.join(resumes_dir, "candidate_03_junior_ml_intern.pdf")
        c = canvas.Canvas(pdf_path, pagesize=letter)
        c.setFont("Helvetica-Bold", 16)
        c.drawString(50, 750, "David Kim — Junior Machine Learning Developer")
        c.setFont("Helvetica", 10)
        c.drawString(50, 735, "Email: dkim.intern@outlook.com | Phone: (555) 901-2345 | GitHub: github.com/dkim-ml")
        c.setFont("Helvetica-Bold", 12)
        c.drawString(50, 705, "Summary:")
        c.setFont("Helvetica", 10)
        c.drawString(50, 690, "Enthusiastic Junior ML developer with 1.0 year of experience in machine learning algorithms,")
        c.drawString(50, 678, "data cleaning, model deployment, and Python scripting.")
        c.setFont("Helvetica-Bold", 12)
        c.drawString(50, 648, "Technical Skills:")
        c.setFont("Helvetica", 10)
        c.drawString(50, 633, "Python, NumPy, Pandas, Matplotlib, Scikit-learn, Basic PyTorch, Docker, Git, Linux")
        c.setFont("Helvetica-Bold", 12)
        c.drawString(50, 603, "Projects:")
        c.setFont("Helvetica", 10)
        c.drawString(50, 588, "- Image Classification Prototype: Trained CNN in PyTorch on CIFAR-10.")
        c.drawString(50, 576, "- House Price Prediction: Built regression model using Scikit-learn and Pandas.")
        c.setFont("Helvetica-Bold", 12)
        c.drawString(50, 546, "Education:")
        c.setFont("Helvetica", 10)
        c.drawString(50, 531, "Bachelor's Degree in Computer Engineering (2024)")
        c.save()
    except Exception:
        # Fallback to TXT if reportlab fails
        txt3 = """David Kim — Junior Machine Learning Developer
Email: dkim.intern@outlook.com | Phone: (555) 901-2345
Summary: Enthusiastic Junior ML developer with 1.0 year of experience in Python, Scikit-learn, PyTorch, Docker, Git.
Education: Bachelor's Degree in Computer Engineering (2024)
"""
        with open(os.path.join(resumes_dir, "candidate_03_junior_ml_intern.txt"), "w", encoding="utf-8") as f:
            f.write(txt3)

    # 4. Lead Full Stack Software Engineer (.docx)
    doc4 = docx.Document()
    doc4.add_heading("Marcus Vance — Lead Full Stack Engineer", 0)
    doc4.add_paragraph("Email: marcus.vance@techcorp.io | Phone: (555) 432-1098")
    doc4.add_heading("Professional Summary", level=1)
    doc4.add_paragraph("Full Stack Software Engineer with 7.0 years of experience designing robust microservices, cloud web applications, and scalable RESTful APIs.")
    doc4.add_heading("Technical Expertise", level=1)
    doc4.add_paragraph("Languages & Frontend: TypeScript, JavaScript, React, Next.js, HTML5, CSS3, Tailwind CSS.")
    doc4.add_paragraph("Backend & Databases: Node.js, Express, FastAPI, PostgreSQL, MongoDB, Redis, GraphQL, REST API.")
    doc4.add_paragraph("DevOps & Practices: Docker, AWS, CI/CD, Git, GitHub Actions, System Design, Unit Testing.")
    doc4.add_heading("Education", level=1)
    doc4.add_paragraph("Bachelor's Degree in Software Engineering, UC Berkeley (2017)")
    doc4.save(os.path.join(resumes_dir, "candidate_04_fullstack_lead.docx"))

    # 5. Frontend Developer (.txt)
    txt5 = """Emily Watson — Frontend React Developer
Email: emily.watson@webdev.com | Phone: 555-654-3210

Summary:
Frontend Web Developer with 3.0 years of experience crafting interactive user interfaces in React, TypeScript, and modern CSS.

Technical Skills:
- React, TypeScript, JavaScript, Next.js, HTML5, CSS3, Tailwind CSS, Redux, Git, Webpack

Experience:
Frontend Developer at PixelCraft (2021 - Present):
- Built component libraries in React and Tailwind CSS.
- Optimized web application performance and accessibility.

Education:
Bachelor's Degree in Graphic Design & Web Informatics (2021)
"""
    with open(os.path.join(resumes_dir, "candidate_05_frontend_dev.txt"), "w", encoding="utf-8") as f:
        f.write(txt5)

    # 6. DevOps & Cloud Infrastructure Architect (.docx)
    doc6 = docx.Document()
    doc6.add_heading("Robert Kowalski — Principal Cloud DevOps Architect", 0)
    doc6.add_paragraph("Email: r.kowalski@cloudops.net | Phone: (555) 789-0123")
    doc6.add_heading("Summary", level=1)
    doc6.add_paragraph("DevOps & Cloud Infrastructure Architect with 8.0 years of experience managing multi-cloud Kubernetes clusters, Infrastructure as Code, and CI/CD pipelines.")
    doc6.add_heading("Skills", level=1)
    doc6.add_paragraph("Cloud & IaC: AWS, Google Cloud Platform, Terraform, Kubernetes, Helm, Docker, Linux, Bash.")
    doc6.add_paragraph("CI/CD & Monitoring: GitHub Actions, Jenkins, Prometheus, Grafana, Ansible, Git, Python.")
    doc6.add_heading("Education", level=1)
    doc6.add_paragraph("Bachelor's Degree in Information Technology (2016)")
    doc6.save(os.path.join(resumes_dir, "candidate_06_devops_cloud_architect.docx"))

    # 7. BI & Data Analyst (.txt)
    txt7 = """Jessica Martinez — Business Intelligence & Data Analyst
Email: jessica.martinez@analytics.org | Phone: (555) 321-7654

Summary:
Data Analyst with 4.5 years of experience delivering executive dashboards, business intelligence insights, and statistical reporting.

Technical Skills:
- SQL, Tableau, Power BI, Excel, Pandas, Python, Business Intelligence, Data Visualization, Metrics Reporting, A/B Testing

Education:
Bachelor's Degree in Business Analytics (2019)
"""
    with open(os.path.join(resumes_dir, "candidate_07_bi_data_analyst.txt"), "w", encoding="utf-8") as f:
        f.write(txt7)

    # 8. Java Backend Engineer (.txt)
    txt8 = """Vikram Patel — Senior Java Backend Developer
Email: vikram.patel@enterprise.com | Phone: (555) 890-5678

Summary:
Backend Engineer with 5.5 years of experience building high-throughput financial microservices.

Technical Skills:
- Java, Spring Boot, PostgreSQL, MySQL, Redis, Kafka, Apache Kafka, REST API, Docker, Kubernetes, Unit Testing, Git

Education:
Bachelor's Degree in Computer Science (2019)
"""
    with open(os.path.join(resumes_dir, "candidate_08_java_backend_dev.txt"), "w", encoding="utf-8") as f:
        f.write(txt8)

    print(f"[DataGenerator] Sample candidate corpus created across PDF, DOCX, and TXT formats in {resumes_dir}")

def generate_sample_job_descriptions(jd_dir: str):
    """Generates structured Job Descriptions across 3 standard roles."""
    os.makedirs(jd_dir, exist_ok=True)
    
    # 1. Senior Machine Learning Engineer
    jd1 = {
        "job_id": "JOB-ML-001",
        "title": "Senior Machine Learning Engineer (NLP & MLOps)",
        "min_experience_years": 4.0,
        "required_skills": [
            "python", "pytorch", "scikit-learn", "machine learning",
            "natural language processing", "docker", "aws", "model deployment"
        ],
        "preferred_skills": [
            "transformers", "huggingface", "kubernetes", "mlops", "langchain", "sql"
        ],
        "description_text": """We are seeking a Senior Machine Learning Engineer to design and deploy scalable NLP and generative AI models. 
Candidates must have strong hands-on experience with Python, PyTorch, Scikit-learn, and Machine Learning algorithms. 
You will be responsible for end-to-end model deployment, Docker containerization, and AWS cloud infrastructure. 
Preferred experience includes Transformers, HuggingFace, Kubernetes, and automated MLOps pipelines."""
    }
    with open(os.path.join(jd_dir, "job_senior_ml_engineer.json"), "w", encoding="utf-8") as f:
        json.dump(jd1, f, indent=2)

    # 2. Lead Full Stack Software Engineer
    jd2 = {
        "job_id": "JOB-FS-002",
        "title": "Lead Full Stack Software Engineer",
        "min_experience_years": 5.0,
        "required_skills": [
            "typescript", "javascript", "react", "node.js", "postgresql", "rest api", "docker"
        ],
        "preferred_skills": [
            "next.js", "graphql", "redis", "aws", "ci/cd", "system design"
        ],
        "description_text": """Looking for an experienced Lead Full Stack Developer to architect modern cloud web applications. 
Must possess expertise in TypeScript, React, Node.js, REST API design, PostgreSQL databases, and Docker containerization. 
Preferred experience includes Next.js, GraphQL, Redis caching, AWS deployments, and CI/CD pipelines."""
    }
    with open(os.path.join(jd_dir, "job_lead_fullstack_engineer.json"), "w", encoding="utf-8") as f:
        json.dump(jd2, f, indent=2)

    # 3. Cloud DevOps & Infrastructure Specialist
    jd3 = {
        "job_id": "JOB-DO-003",
        "title": "Senior Cloud DevOps & Infrastructure Specialist",
        "min_experience_years": 4.0,
        "required_skills": [
            "aws", "kubernetes", "docker", "terraform", "ci/cd", "linux"
        ],
        "preferred_skills": [
            "python", "prometheus", "grafana", "github actions", "bash", "ansible"
        ],
        "description_text": """Seeking a Senior Cloud DevOps Specialist to lead automated cloud infrastructure and CI/CD operations. 
Required hands-on skills include AWS, Kubernetes cluster management, Docker, Terraform Infrastructure as Code, Linux systems, and CI/CD pipelines. 
Preferred skills include Prometheus monitoring, Grafana dashboards, GitHub Actions, and Python automation scripting."""
    }
    with open(os.path.join(jd_dir, "job_cloud_devops_specialist.json"), "w", encoding="utf-8") as f:
        json.dump(jd3, f, indent=2)

    print(f"[DataGenerator] 3 structured Job Descriptions saved to {jd_dir}")

if __name__ == "__main__":
    base = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
    generate_sample_resumes(os.path.join(base, "data", "resumes"))
    generate_sample_job_descriptions(os.path.join(base, "data", "job_descriptions"))
